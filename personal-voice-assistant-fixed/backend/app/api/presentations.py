import json
from pathlib import Path
from uuid import uuid4
from fastapi import APIRouter, HTTPException, Request
from pydantic import BaseModel, Field
from ..core.model_adapters import ModelError

router=APIRouter(prefix="/api/presentations",tags=["presentations"])

class PresentationRequest(BaseModel):
    topic: str = Field(min_length=3,max_length=3000)
    slide_count: int = Field(default=8,ge=3,le=20)
    audience: str = Field(default="general professional audience",max_length=500)
    style: str = Field(default="clean professional",max_length=300)

def _slides_from_response(text: str, topic: str, count: int):
    try:
        start=text.find("{"); end=text.rfind("}")
        payload=json.loads(text[start:end+1])
        slides=payload.get("slides",[])
        clean=[]
        for item in slides[:count]:
            title=str(item.get("title","")).strip()[:160]
            bullets=item.get("bullets",[])
            if title:
                clean.append({"title":title,"bullets":[str(x)[:500] for x in bullets[:6]]})
        if clean: return clean
    except Exception: pass
    return [{"title":topic[:100],"bullets":["Review the topic and objectives","Replace this placeholder with a verified key point"]}]

@router.post("/generate")
async def generate_presentation(body: PresentationRequest, request: Request):
    prompt=("Create an editable PowerPoint slide outline. Return ONLY valid JSON with shape "
            '{"title":"...", "slides":[{"title":"...", "bullets":["..."]}]}. '
            f"Create {body.slide_count} slides for audience {body.audience!r}, visual style {body.style!r}, topic {body.topic!r}. "
            "Use concise, factual bullets. Do not invent statistics or citations. Include an introduction and conclusion.")
    try: result=await request.app.state.models.chat([{"role":"user","content":prompt}])
    except ModelError as exc: raise HTTPException(503,str(exc)) from exc
    slides=_slides_from_response(result["content"],body.topic,body.slide_count)
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise HTTPException(503,"PowerPoint generation requires python-pptx. Install it with `pip install python-pptx`.") from exc
    settings=request.app.state.settings
    outdir=Path(settings.generated_files_dir).expanduser().resolve()
    outdir.mkdir(parents=True,exist_ok=True)
    filename=f"presentation-{uuid4()}.pptx"
    outpath=outdir/filename
    prs=Presentation()
    prs.core_properties.title=body.topic
    title_slide=prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text=body.topic
    if len(title_slide.placeholders)>1: title_slide.placeholders[1].text=f"{body.audience} · Generated with AVA"
    for item in slides:
        slide=prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text=item["title"]
        tf=slide.placeholders[1].text_frame
        tf.clear()
        bullets=item["bullets"] or ["Add verified content for this slide."]
        for i,bullet in enumerate(bullets):
            para=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            para.text=bullet; para.level=0; para.font.size=Pt(22)
    prs.save(outpath)
    return {"title":body.topic,"slides":len(slides)+1,"file_name":filename,"download_url":f"/api/presentations/files/{filename}","model":result.get("model"),"editable":True}

@router.get("/files/{filename}")
def download_presentation(filename: str, request: Request):
    from fastapi.responses import FileResponse
    if "/" in filename or "\\" in filename or not filename.endswith(".pptx"):
        raise HTTPException(400,"Invalid filename.")
    path=(Path(request.app.state.settings.generated_files_dir).expanduser().resolve()/filename).resolve()
    base=Path(request.app.state.settings.generated_files_dir).expanduser().resolve()
    if path.parent != base or not path.is_file(): raise HTTPException(404,"Presentation not found.")
    return FileResponse(path,filename=filename,media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")


@router.post("/files/{filename}/export-pdf")
def export_presentation_pdf(filename: str, request: Request):
    import shutil, subprocess
    from fastapi.responses import FileResponse
    if "/" in filename or "\\" in filename or not filename.endswith(".pptx"):
        raise HTTPException(400,"Invalid filename.")
    base=Path(request.app.state.settings.generated_files_dir).expanduser().resolve()
    source=(base/filename).resolve()
    if source.parent!=base or not source.is_file(): raise HTTPException(404,"Presentation not found.")
    converter=shutil.which("libreoffice") or shutil.which("soffice")
    if not converter: raise HTTPException(503,"PDF export requires LibreOffice installed on this machine.")
    try:
        subprocess.run([converter,"--headless","--convert-to","pdf","--outdir",str(base),str(source)],
                       capture_output=True,text=True,timeout=60,check=True)
    except subprocess.TimeoutExpired as exc: raise HTTPException(408,"PDF conversion timed out.") from exc
    except subprocess.CalledProcessError as exc: raise HTTPException(502,f"PDF conversion failed: {exc.stderr[-1000:]}") from exc
    pdf=(base/(source.stem+".pdf")).resolve()
    if pdf.parent!=base or not pdf.is_file(): raise HTTPException(502,"PDF was not produced.")
    return FileResponse(pdf,filename=pdf.name,media_type="application/pdf")
